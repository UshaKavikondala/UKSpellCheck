import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .agents.aggregator_agent import KIND_PRIORITY
from .state import PipelineState

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
KIND_FILL = {
    "spelling": "F8CBAD",
    "foreign_word": "FFE699",
    "grammar": "BDD7EE",
    "punctuation": "C6E0B4",
    "language": "D9D9D9",
}
ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _put(ws, row: int, col: int, value):
    """Write a cell. Text is always stored as text, so page content starting with '=' is never run as a formula."""
    cell = ws.cell(row=row, column=col)
    if isinstance(value, str):
        cell.value = ILLEGAL_XML.sub("", value)
        cell.data_type = "s"
    else:
        cell.value = value
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    return cell


def _header(ws, titles: list[str]) -> None:
    for col, title in enumerate(titles, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    ws.freeze_panes = "A2"


def _widths(ws, widths: list[int]) -> None:
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _status(state: PipelineState) -> str:
    if not state.errors:
        return "OK"
    return "PARTIAL" if state.final_issues else "FAILED"


def write_excel(states: list[PipelineState], path: Path) -> Path:
    wb = Workbook()

    # Sheet 1: one row per URL
    summary = wb.active
    summary.title = "Summary"
    _header(summary, ["URL", "Status", "HTTP", "Page language", "Blocks", "Chunks checked",
                      *[k for k in KIND_PRIORITY], "Total issues", "API calls", "Tokens",
                      "Seconds", "Errors"])
    for r, s in enumerate(states, start=2):
        counts = {k: sum(1 for i in s.final_issues if i.kind == k) for k in KIND_PRIORITY}
        row = [s.url or "(text input)", _status(s), s.http_status, s.page_language,
               len(s.blocks), len(s.chunks), *counts.values(), len(s.final_issues),
               s.api_calls, s.total_tokens, round(s.elapsed_seconds, 1), "; ".join(s.errors)]
        for c, value in enumerate(row, start=1):
            _put(summary, r, c, value)
    _widths(summary, [60, 10, 7, 14, 8, 10, 10, 12, 10, 12, 10, 12, 10, 10, 9, 60])
    summary.auto_filter.ref = summary.dimensions

    # Sheet 2: one row per issue, filterable by URL and kind
    issues = wb.create_sheet("Issues")
    _header(issues, ["URL", "Kind", "Text", "Suggestion", "Reason", "Chunk", "Count"])
    r = 2
    for s in states:
        for i in s.final_issues:
            row = [s.url or "(text input)", i.kind, i.text, i.suggestion, i.reason,
                   i.chunk_index if i.chunk_index >= 0 else "", i.count]
            for c, value in enumerate(row, start=1):
                cell = _put(issues, r, c, value)
                if c == 2:
                    cell.fill = PatternFill("solid", fgColor=KIND_FILL.get(i.kind, "FFFFFF"))
            r += 1
    _widths(issues, [60, 14, 45, 45, 60, 8, 8])
    issues.auto_filter.ref = issues.dimensions

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def default_report_path() -> Path:
    stamp = datetime.now().strftime("%Y%m%dT%H%M%S")
    return Path(__file__).resolve().parent / "reports" / f"spellcheck_report_{stamp}.xlsx"
